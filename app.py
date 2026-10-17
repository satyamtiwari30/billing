import json
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import streamlit as st
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

# -------------------- Setup --------------------
st.set_page_config(page_title="ABCD COMPUTERS Billing", page_icon="🧾", layout="wide")
st.title("🧾 ABCD COMPUTERS — Billing Invoice")

BASE_DIR = Path(__file__).resolve().parent
SETTINGS_FILE = BASE_DIR / "company_settings.json"
EXCEL_FILE = BASE_DIR / "invoice_records.xlsx"
INVOICE_DIR = BASE_DIR / "invoices"
INVOICE_DIR.mkdir(exist_ok=True)

SUMMARY_HEADERS = [
    "Invoice No", "Date", "Company", "Address", "Customer", "Customer Address", "Total", "PDF File"
]
ITEM_HEADERS = [
    "Invoice No", "Date", "Customer", "Product", "Quantity", "Price", "Amount"
]


def load_settings():
    if SETTINGS_FILE.exists():
        try:
            data = json.loads(SETTINGS_FILE.read_text(encoding="utf-8"))
            return {"company": data.get("company", ""),
                    "address": data.get("address", ""),
                    "logo": data.get("logo", "")}
        except (json.JSONDecodeError, OSError):
            pass
    return {"company": "", "address": "", "logo": ""}


def initialize_excel():
    if not EXCEL_FILE.exists():
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Invoices"
        ws.append(SUMMARY_HEADERS)
        items = wb.create_sheet("Invoice Items")
        items.append(ITEM_HEADERS)
        wb.save(EXCEL_FILE)
        wb.close()
    else:
        # Add the item sheet if this is an older workbook from the simple version.
        wb = openpyxl.load_workbook(EXCEL_FILE)
        if "Invoices" not in wb.sheetnames:
            ws = wb.active
            ws.title = "Invoices"
            if ws.max_row == 1 and ws.max_column == 1 and ws["A1"].value is None:
                ws.append(SUMMARY_HEADERS)
        if "Invoice Items" not in wb.sheetnames:
            items = wb.create_sheet("Invoice Items")
            items.append(ITEM_HEADERS)
        wb.save(EXCEL_FILE)
        wb.close()


def format_sheet(ws):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C00000")
        cell.alignment = Alignment(horizontal="center")
    for cells in ws.columns:
        letter = get_column_letter(cells[0].column)
        longest = max((len(str(c.value)) for c in cells if c.value is not None), default=12)
        ws.column_dimensions[letter].width = min(max(longest + 2, 12), 38)


def next_invoice_number():
    initialize_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    ws = wb["Invoices"]
    nums = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        val = row[0] if row else None
        if val:
            try:
                nums.append(int(str(val).split("-")[-1]))
            except ValueError:
                continue
    wb.close()
    return f"INV-{max(nums, default=0) + 1:04d}"


def save_to_excel(invoice, items, pdf_name):
    initialize_excel()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    summary = wb["Invoices"]
    item_sheet = wb["Invoice Items"]

    existing = {str(row[0]) for row in summary.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}
    if invoice["invoice_no"] in existing:
        wb.close()
        raise ValueError("This invoice number already exists. Please refresh and try again.")

    summary.append([
        invoice["invoice_no"], invoice["invoice_date"], invoice["company"],
        invoice["address"], invoice["customer"], invoice["customer_address"], invoice["total"], pdf_name
    ])
    for item in items:
        item_sheet.append([
            invoice["invoice_no"], invoice["invoice_date"], invoice["customer"],
            item["Product Name"], item["Quantity"], item["Price"], item["Amount"]
        ])
    format_sheet(summary)
    format_sheet(item_sheet)
    wb.save(EXCEL_FILE)
    wb.close()


def make_pdf(invoice, items, logo_bytes=None):
    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4
    left, right = 45, width - 45
    y = height - 45

    # Logo, if one has been saved/uploaded
    if logo_bytes:
        try:
            pdf.drawImage(ImageReader(BytesIO(logo_bytes)), left, y - 55,
                          width=72, height=50, preserveAspectRatio=True,
                          anchor="c", mask="auto")
        except Exception:
            pass

    pdf.setFillColorRGB(0.75, 0.0, 0.0)
    pdf.setFont("Helvetica-Bold", 19)
    pdf.drawString(left + (85 if logo_bytes else 0), y - 20, invoice["company"][:55])
    pdf.setFillColorRGB(0, 0, 0)
    pdf.setFont("Helvetica", 9)
    address_y = y - 42
    for line in invoice["address"].splitlines()[:4]:
        pdf.drawString(left + (85 if logo_bytes else 0), address_y, line[:85])
        address_y -= 12

    y = min(address_y, y - 75) - 20
    pdf.setStrokeColorRGB(0.75, 0, 0)
    pdf.setLineWidth(1.5)
    pdf.line(left, y, right, y)
    y -= 28
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawString(left, y, "INVOICE")
    y -= 25
    pdf.setFont("Helvetica", 10)
    pdf.drawString(left, y, f"Invoice No: {invoice['invoice_no']}")
    pdf.drawString(left + 250, y, f"Date: {invoice['invoice_date']}")
    y -= 18
    pdf.drawString(left, y, f"Bill To: {invoice['customer'][:70]}")
    y -= 16
    pdf.setFont("Helvetica", 9)
    customer_address_lines = invoice.get("customer_address", "").splitlines() or [""]
    for line in customer_address_lines[:3]:
        pdf.drawString(left, y, f"Address: {line[:90]}")
        y -= 13
    y -= 15

    # Product table header
    pdf.setFillColorRGB(0.75, 0, 0)
    pdf.rect(left, y - 5, right - left, 22, fill=1, stroke=0)
    pdf.setFillColorRGB(1, 1, 1)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(left + 8, y + 3, "Product / Service")
    pdf.drawRightString(left + 325, y + 3, "Qty")
    pdf.drawRightString(left + 410, y + 3, "Price (Rs.)")
    pdf.drawRightString(right - 8, y + 3, "Amount (Rs.)")
    y -= 24

    pdf.setFillColorRGB(0, 0, 0)
    pdf.setFont("Helvetica", 9)
    for item in items:
        if y < 95:
            pdf.showPage()
            y = height - 50
            pdf.setFont("Helvetica", 9)
        name = str(item["Product Name"]).replace("\n", " ")[:48]
        pdf.drawString(left + 8, y, name)
        pdf.drawRightString(left + 325, y, f"{item['Quantity']:g}")
        pdf.drawRightString(left + 410, y, f"{item['Price']:,.2f}")
        pdf.drawRightString(right - 8, y, f"{item['Amount']:,.2f}")
        pdf.setStrokeColorRGB(0.82, 0.82, 0.82)
        pdf.line(left, y - 6, right, y - 6)
        y -= 23

    y -= 15
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawRightString(right - 8, y, f"Grand Total: Rs. {invoice['total']:,.2f}")
    y -= 35
    pdf.setFont("Helvetica", 9)
    pdf.drawString(left, y, "Thank you for your business!")
    pdf.save()
    return buffer.getvalue()


settings = load_settings()

# -------------------- Company settings --------------------
st.header("Company Settings")
company = st.text_input("Company Name", value=settings["company"])
address = st.text_area("Company Address", value=settings["address"])
logo_upload = st.file_uploader("Upload Company Logo", type=["png", "jpg", "jpeg"])

saved_logo_path = BASE_DIR / settings["logo"] if settings.get("logo") else None
logo_bytes = logo_upload.getvalue() if logo_upload else None
if logo_bytes:
    st.image(logo_bytes, width=150)
elif saved_logo_path and saved_logo_path.exists():
    logo_bytes = saved_logo_path.read_bytes()
    st.image(logo_bytes, width=150, caption="Saved company logo")

if st.button("Save Company Details"):
    settings["company"] = company.strip()
    settings["address"] = address.strip()
    if logo_upload:
        logo_path = BASE_DIR / "company_logo.png"
        logo_path.write_bytes(logo_upload.getvalue())
        settings["logo"] = logo_path.name
    SETTINGS_FILE.write_text(json.dumps(settings, indent=2), encoding="utf-8")
    st.success("Company details saved. They will be prefilled next time you open the app.")
    st.rerun()

st.divider()

# -------------------- Customer details --------------------
st.header("Customer Details")
customer = st.text_input("Customer Name")
customer_address = st.text_area("Customer Address", placeholder="Enter customer billing address")
invoice_date = st.date_input("Invoice Date", value=date.today())

# -------------------- Product details --------------------
st.header("Product Details")
if "products_table" not in st.session_state:
    st.session_state.products_table = pd.DataFrame([
        {"Product Name": "", "Quantity": 1, "Price": 0.0}
    ])

products = st.data_editor(
    st.session_state.products_table,
    num_rows="dynamic",
    use_container_width=True,
    hide_index=True,
    key="product_editor",
    column_config={
        "Product Name": st.column_config.TextColumn("Product Name", required=True),
        "Quantity": st.column_config.NumberColumn("Quantity", min_value=1, step=1, required=True),
        "Price": st.column_config.NumberColumn("Price per Item (Rs.)", min_value=0.0, step=100.0, required=True, format="%.2f"),
    },
)

# Work on a copy, discard blank rows, and safely validate numbers.
products = products.copy()
products["Product Name"] = products["Product Name"].fillna("").astype(str).str.strip()
products["Quantity"] = pd.to_numeric(products["Quantity"], errors="coerce")
products["Price"] = pd.to_numeric(products["Price"], errors="coerce")
valid_products = products[products["Product Name"] != ""].copy()
valid_products = valid_products.dropna(subset=["Quantity", "Price"])
valid_products = valid_products[(valid_products["Quantity"] > 0) & (valid_products["Price"] >= 0)]
valid_products["Amount"] = (valid_products["Quantity"] * valid_products["Price"]).round(2)

total = round(float(valid_products["Amount"].sum()), 2) if not valid_products.empty else 0.0
if not valid_products.empty:
    st.dataframe(valid_products, use_container_width=True, hide_index=True)
st.metric("Grand Total", f"₹{total:,.2f}")

# -------------------- Generate invoice --------------------
if st.button("Generate Invoice", type="primary"):
    if not company.strip() or not address.strip() or not customer.strip() or not customer_address.strip():
        st.error("Please fill in company name, company address, customer name, and customer address.")
    elif valid_products.empty:
        st.error("Please add at least one product with a valid quantity and price.")
    elif (valid_products["Quantity"] <= 0).any():
        st.error("Quantity must be greater than zero for every product.")
    else:
        try:
            initialize_excel()
            invoice_id = next_invoice_number()
            pdf_name = f"{invoice_id}.pdf"
            invoice = {
                "invoice_no": invoice_id,
                "invoice_date": invoice_date.strftime("%d-%m-%Y"),
                "company": company.strip(),
                "address": address.strip(),
                "customer": customer.strip(),
                "customer_address": customer_address.strip(),
                "total": total,
            }
            item_records = valid_products[["Product Name", "Quantity", "Price", "Amount"]].to_dict("records")
            pdf_bytes = make_pdf(invoice, item_records, logo_bytes)
            (INVOICE_DIR / pdf_name).write_bytes(pdf_bytes)
            save_to_excel(invoice, item_records, pdf_name)
            st.success(f"Invoice {invoice_id} generated and saved to PDF and Excel.")
            st.download_button("Download PDF Invoice", data=pdf_bytes, file_name=pdf_name, mime="application/pdf")
        except Exception as exc:
            st.error(f"Could not generate invoice: {exc}")

# -------------------- Invoice history / Excel download --------------------
st.divider()
st.header("Invoice Records")
if st.button("Refresh Invoice Records"):
    st.rerun()
try:
    initialize_excel()
    history = pd.read_excel(EXCEL_FILE, sheet_name="Invoices")
    if not history.empty:
        st.dataframe(history, use_container_width=True, hide_index=True)

        st.subheader("Delete an Old Invoice")
        st.warning("Deleting an invoice removes its row from Invoices, its items from Invoice Items, and its saved PDF. This cannot be undone. Download a backup first if needed.")
        invoice_options = history["Invoice No"].dropna().astype(str).tolist()
        if invoice_options:
            selected_invoice = st.selectbox("Choose invoice to delete", options=invoice_options, key="delete_invoice_select")
            confirm_delete = st.checkbox(f"I confirm that I want to permanently delete {selected_invoice}", key="confirm_delete_invoice")
            if st.button("Delete Selected Invoice", type="secondary", disabled=not confirm_delete):
                try:
                    wb = openpyxl.load_workbook(EXCEL_FILE)
                    for sheet_name in ("Invoices", "Invoice Items"):
                        if sheet_name in wb.sheetnames:
                            ws = wb[sheet_name]
                            header = [cell.value for cell in ws[1]]
                            try:
                                invoice_col = header.index("Invoice No") + 1
                            except ValueError:
                                continue
                            for row_idx in range(ws.max_row, 1, -1):
                                if str(ws.cell(row=row_idx, column=invoice_col).value) == selected_invoice:
                                    ws.delete_rows(row_idx, 1)
                    wb.save(EXCEL_FILE)
                    wb.close()

                    # Delete only the PDF filename recorded for this selected invoice.
                    matching = history[history["Invoice No"].astype(str) == selected_invoice]
                    if not matching.empty and "PDF File" in matching.columns:
                        pdf_value = matching.iloc[0]["PDF File"]
                        if pd.notna(pdf_value) and str(pdf_value).strip():
                            pdf_path = (INVOICE_DIR / Path(str(pdf_value)).name).resolve()
                            if pdf_path.parent == INVOICE_DIR.resolve() and pdf_path.exists():
                                pdf_path.unlink()
                    st.success(f"Invoice {selected_invoice} and its related records have been deleted.")
                    st.rerun()
                except Exception as exc:
                    st.error(f"Could not delete invoice {selected_invoice}: {exc}")
    else:
        st.info("No invoices have been saved yet.")
    st.download_button(
        "Download Excel Records",
        data=EXCEL_FILE.read_bytes(),
        file_name="invoice_records.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
except Exception as exc:
    st.warning(f"Could not load invoice records: {exc}")
